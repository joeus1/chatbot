"""Pure chat logic for the Streamlit app.

Everything here is importable without Streamlit so it can be unit-tested
with no network and no UI. The app (`streamlit_app.py`) owns presentation
and session state; this module owns message shaping and error mapping.
"""

import base64
import csv
import io
import logging
from pathlib import Path

from openai import (
    APIConnectionError,
    AuthenticationError,
    BadRequestError,
    RateLimitError,
    UnprocessableEntityError,
)

logger = logging.getLogger(__name__)

# User-facing copy for API failures. Raw exception text never reaches the
# page: it can contain request details and, for auth errors, hints about
# the key.
AUTH_ERROR_MESSAGE = (
    "OpenAI rejected the configured API key. "
    "Check OPENAI_API_KEY in your Streamlit secrets and try again."
)
RATE_LIMIT_MESSAGE = (
    "The assistant is receiving too many requests right now. "
    "Wait a moment and send your message again."
)
CONNECTION_ERROR_MESSAGE = (
    "Could not reach OpenAI. Check your network connection and try again."
)
GENERIC_ERROR_MESSAGE = (
    "Something went wrong while generating a response. "
    "Your message was kept; try sending it again."
)
NON_RETRYABLE_MESSAGE = (
    "OpenAI could not process that message, so it was removed from the "
    "conversation. Try a shorter or rephrased version - and use Clear "
    "conversation if it keeps happening, since a long history can cause "
    "this too."
)
EMPTY_RESPONSE_MESSAGE = (
    "The assistant returned an empty reply. Nothing failed - ask again or "
    "rephrase if you were expecting an answer."
)

# Schedule uploads. Text formats are read and sent as text; images go to the
# model as images. Anything else is refused with a message rather than a
# traceback, and the bounds below keep one upload from blowing the context
# window or the request size.
TEXT_UPLOAD_TYPES = ("csv", "tsv", "txt", "md")
SHEET_UPLOAD_TYPES = ("xlsx",)
IMAGE_UPLOAD_TYPES = ("png", "jpg", "jpeg", "webp")
UPLOAD_TYPES = TEXT_UPLOAD_TYPES + SHEET_UPLOAD_TYPES + IMAGE_UPLOAD_TYPES
MAX_UPLOAD_TEXT_CHARS = 40_000
# Byte ceilings checked before anything is decoded or parsed, so an oversized
# file is refused for the cost of len(), not a full openpyxl load. Text is
# capped at four bytes per allowed character (the widest UTF-8 encoding); a
# workbook at 4 MB is already far past anything that renders to 40k chars.
MAX_UPLOAD_TEXT_BYTES = 4 * MAX_UPLOAD_TEXT_CHARS
MAX_UPLOAD_SHEET_BYTES = 4 * 1024 * 1024
MAX_UPLOAD_IMAGE_BYTES = 8 * 1024 * 1024
UPLOAD_INSTRUCTION = (
    "Here is the schedule to review. Read it back per section 3, then ask "
    "for the intake facts and the preset questions."
)
UNSUPPORTED_UPLOAD_MESSAGE = (
    "That file type is not supported. Upload the schedule as CSV, TSV, TXT, "
    "XLSX, PNG, JPG or WEBP, or paste it as text."
)
UPLOAD_TOO_LARGE_MESSAGE = (
    "That file is too large to review in one go. Trim it to the period you "
    "want reviewed, or upload one week at a time."
)
EMPTY_UPLOAD_MESSAGE = "That file is empty. Check the export and upload it again."
# Typed and pasted messages. The 40k-character upload bound above does not
# reach the chat box, so without this a paste could be as large as the model's
# context and be resent on every turn of the window. Same bound as an upload:
# a schedule that fits one has room for the questions and answers around it.
MAX_CHAT_MESSAGE_CHARS = MAX_UPLOAD_TEXT_CHARS
MESSAGE_TOO_LONG_MESSAGE = (
    "That message is too long to review in one go (the limit is "
    f"{MAX_CHAT_MESSAGE_CHARS:,} characters). Trim it to the period you want "
    "reviewed, or send one week at a time."
)
# One review is a read-back, the intake facts, seven answers per person and a
# few corrections; a busy store is well under this. It bounds what one session
# can spend, since the app has no login. Clear conversation starts a new
# budget, so this slows a runaway loop rather than stopping a determined
# abuser: the OpenAI key's own spend limit is the hard stop.
MAX_USER_TURNS_PER_SESSION = 100
SESSION_TURN_LIMIT_MESSAGE = (
    "This conversation has reached its limit of "
    f"{MAX_USER_TURNS_PER_SESSION} messages. Use Clear conversation in the "
    "sidebar to start a new review."
)
UNREADABLE_SHEET_MESSAGE = (
    "That workbook could not be read. It may be an older .xls renamed, "
    "password-protected, or damaged. Re-export it as .xlsx or CSV and try again."
)


class UploadError(ValueError):
    """An upload the app should refuse with the message it carries."""


class MessageError(ValueError):
    """A typed message the app should refuse with the message it carries."""


def check_chat_message(text):
    """Refuse a typed message longer than `MAX_CHAT_MESSAGE_CHARS`.

    The chat box's own `max_chars` is enforced in the browser, so a client
    that talks to the server directly skips it; this is the check that holds.
    """
    if len(text) > MAX_CHAT_MESSAGE_CHARS:
        raise MessageError(MESSAGE_TOO_LONG_MESSAGE)


def session_turn_limit_reached(history, limit=MAX_USER_TURNS_PER_SESSION):
    """Whether `history` already holds `limit` user turns.

    Counts what the transcript holds, not what was ever sent: a turn dropped
    as non-retryable was refused by the API, so it does not use up the budget.
    """
    return sum(1 for m in history if m["role"] == "user") >= limit


def append_message(history, role, content, display=None):
    """Append a validated chat turn to `history` in place.

    `content` is what the API receives: a non-empty string, or a list of
    content parts for a turn that carries an upload. A list needs `display`,
    the text the page shows in its place, so a schedule's raw cells (which
    may hold full names) are never rendered back into the transcript.
    """
    if role not in ("user", "assistant"):
        raise ValueError(f"unsupported role: {role!r}")
    if isinstance(content, list):
        if not content:
            raise ValueError("content parts must be a non-empty list")
        if not isinstance(display, str) or not display.strip():
            raise ValueError("a content-part turn needs display text")
        history.append({"role": role, "content": content, "display": display})
        return
    if not isinstance(content, str) or not content.strip():
        raise ValueError("message content must be a non-empty string")
    history.append({"role": role, "content": content})


def display_text(message):
    """The text the page renders for a turn."""
    return message.get("display", message["content"])


def drop_last_message(history, role):
    """Remove the final turn in place if it has `role`.

    Returns True when a turn was removed, so callers can tell the difference
    between undoing their own append and finding nothing to undo.
    """
    if history and history[-1]["role"] == role:
        history.pop()
        return True
    return False


def should_keep_turn(exc):
    """Whether the user's turn should stay in history after `exc`.

    Auth, rate-limit, connection and server errors are all worth retrying with
    the same message unchanged, so the turn stays and the user can resend once
    the key, the quota or the network is fixed.

    A 400 or 422 is caused by the request content itself - an over-long context
    or rejected content. Keeping that turn would make every later message fail
    identically, and the bounded window can never evict it, so the chat wedges
    until the whole conversation is cleared.
    """
    return not isinstance(exc, (BadRequestError, UnprocessableEntityError))


def build_api_messages(history, system_prompt, max_turns):
    """Return the payload sent to the API: system prompt + last `max_turns` turns.

    The full transcript stays in session state for display; only a bounded
    window goes to the API so long chats don't grow cost quadratically or
    overflow the context window.
    """
    if max_turns < 1:
        raise ValueError("max_turns must be at least 1")
    bounded = history[-max_turns:]
    # The schedule itself is the thing under review, and a review runs to many
    # more turns than the window holds: read-back, intake facts, seven answers
    # per person, corrections. If the window has scrolled past the most recent
    # upload, carry that one turn forward ahead of the window so the model is
    # never asked about a sheet it can no longer see. Only the latest upload
    # is pinned; an earlier sheet the manager replaced stays evicted.
    pinned = []
    if not any(is_upload_turn(m) for m in bounded):
        for m in reversed(history[: len(history) - len(bounded)]):
            if is_upload_turn(m):
                pinned = [m]
                break
    return [{"role": "system", "content": system_prompt}] + [
        {"role": m["role"], "content": m["content"]} for m in pinned + bounded
    ]


def is_upload_turn(message):
    """Whether a stored turn carries an upload (content parts, not a string)."""
    return isinstance(message.get("content"), list)


def friendly_error(exc):
    """Map an OpenAI exception to a user-facing message.

    Never includes `str(exc)`; callers may log the exception server-side.
    """
    if isinstance(exc, AuthenticationError):
        return AUTH_ERROR_MESSAGE
    if isinstance(exc, RateLimitError):
        return RATE_LIMIT_MESSAGE
    # APITimeoutError subclasses APIConnectionError, so one branch covers both.
    if isinstance(exc, APIConnectionError):
        return CONNECTION_ERROR_MESSAGE
    # Checked after the config-caused errors above, which keep their own copy.
    if not should_keep_turn(exc):
        return NON_RETRYABLE_MESSAGE
    return GENERIC_ERROR_MESSAGE


def load_system_prompt(path):
    """Read the system prompt from `path`; refuse an empty file.

    The prompt is a document, not a constant, so it lives in a file people can
    read and diff. An empty or missing file would silently ship the model with
    no instructions, so both are errors.
    """
    text = Path(path).read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError(f"system prompt at {path} is empty")
    return text


def _sheet_to_text(data):
    """Render every worksheet of an .xlsx as tab-separated rows.

    Imported lazily so the module stays importable without openpyxl in
    environments that only run the text path.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        out = io.StringIO()
        writer = csv.writer(out, delimiter="\t", lineterminator="\n")
        for sheet in workbook.worksheets:
            heading_written = False
            for row in sheet.iter_rows(values_only=True):
                if all(cell is None for cell in row):
                    continue
                if not heading_written:
                    # Written on the first real row, so a sheet with no cells
                    # contributes nothing and an empty workbook stays empty.
                    out.write(f"## Sheet: {sheet.title}\n")
                    heading_written = True
                writer.writerow("" if cell is None else str(cell) for cell in row)
                if out.tell() > MAX_UPLOAD_TEXT_CHARS:
                    # Past the bound already; the caller refuses it, so there
                    # is no point rendering the rest of the workbook.
                    return out.getvalue()
        return out.getvalue()
    finally:
        workbook.close()


def upload_to_turn(name, data):
    """Turn an uploaded schedule into (api_content, display) for a user turn.

    Text and spreadsheets become a single text block after the instruction;
    images become an image content part. Raises UploadError with user-facing
    copy for anything the app should refuse. The display string names the file
    only: the cells themselves may carry full names, and the transcript on
    screen must not.
    """
    suffix = Path(name).suffix.lower().lstrip(".")
    if suffix not in UPLOAD_TYPES:
        raise UploadError(UNSUPPORTED_UPLOAD_MESSAGE)
    if not data:
        raise UploadError(EMPTY_UPLOAD_MESSAGE)
    display = f"Uploaded schedule: `{Path(name).name}`"

    if suffix in IMAGE_UPLOAD_TYPES:
        if len(data) > MAX_UPLOAD_IMAGE_BYTES:
            raise UploadError(UPLOAD_TOO_LARGE_MESSAGE)
        mime = "image/jpeg" if suffix in ("jpg", "jpeg") else f"image/{suffix}"
        encoded = base64.b64encode(data).decode("ascii")
        content = [
            {"type": "text", "text": UPLOAD_INSTRUCTION},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{encoded}"}},
        ]
        return content, display

    if suffix in SHEET_UPLOAD_TYPES:
        if len(data) > MAX_UPLOAD_SHEET_BYTES:
            raise UploadError(UPLOAD_TOO_LARGE_MESSAGE)
        try:
            text = _sheet_to_text(data)
        except ImportError:
            # A deployment without openpyxl is an operator's defect, not the
            # manager's file; surfacing it as an upload error would send them
            # off to re-export a workbook that was never read.
            raise
        except Exception as exc:  # noqa: BLE001 - any parse failure is the user's file
            logger.warning("workbook upload could not be parsed: %s", type(exc).__name__)
            raise UploadError(UNREADABLE_SHEET_MESSAGE) from exc
    else:
        if len(data) > MAX_UPLOAD_TEXT_BYTES:
            raise UploadError(UPLOAD_TOO_LARGE_MESSAGE)
        text = data.decode("utf-8", errors="replace")
    if not text.strip():
        raise UploadError(EMPTY_UPLOAD_MESSAGE)
    if len(text) > MAX_UPLOAD_TEXT_CHARS:
        raise UploadError(UPLOAD_TOO_LARGE_MESSAGE)
    content = [
        {"type": "text", "text": f"{UPLOAD_INSTRUCTION}\n\nFile: {Path(name).name}\n\n{text}"}
    ]
    return content, display
